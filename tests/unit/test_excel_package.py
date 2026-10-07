"""Excel package pre-checks and data-only parsing (tasks 8.1-8.3): XLS-01, XLS-02, XLS-03.

Pure tests over generated workbooks: no storage, and every test that opens a workbook runs
with outbound sockets blocked (``no_network``), so "no network request is made" is proved by
construction rather than assumed.
"""

from __future__ import annotations

import io
import socket
import zipfile
from typing import Any

import pytest
from finplan_contracts.canonical import canonicalize, sha256_hex

from finplan_platform.core.config import load_config
from finplan_platform.excel import ImportRejected, TemplateSpec, build_workbook, content_from_workbook, import_workbook_bytes
from finplan_platform.excel.package import PackageLimits, open_package
from finplan_platform.excel.template import TEMPLATE_VERSION
from tests.unit.excel_support import SAMPLE_CONTENT, SAMPLE_META, clean, craft, rezip, sheet_xml, unzip

LIMITS = load_config("beta").limits
pytestmark = pytest.mark.usefixtures("no_network")


def rejected(data: bytes, reason: str, *, limits: Any = None, spec: TemplateSpec | None = None) -> ImportRejected:
    with pytest.raises(ImportRejected) as ei:
        import_workbook_bytes(data, limits or LIMITS, spec=spec)
    assert ei.value.code == "VALIDATION_FAILED"
    assert ei.value.details["reason"] == reason, (ei.value.details, ei.value.message)
    return ei.value


# ===================================================================== XLS-01 export template
def test_export_round_trips_exactly_XLS_01() -> None:
    data = clean()
    parsed = import_workbook_bytes(data, LIMITS)
    assert parsed.logical["template_version"] == TEMPLATE_VERSION
    assert parsed.metadata == {k: SAMPLE_META[k] for k in ("plan_id", "base_plan_version_id", "base_checksum", "expected_revision", "exported_at")}
    content = content_from_workbook(SAMPLE_CONTENT, parsed)
    assert content == SAMPLE_CONTENT
    assert sha256_hex(canonicalize(content)) == sha256_hex(canonicalize(SAMPLE_CONTENT))
    # deterministic bytes, no macros/formulas/links/properties
    assert build_workbook(SAMPLE_CONTENT, SAMPLE_META) == data
    parts = unzip(data)
    assert sorted(parts) == sorted(["[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/styles.xml", "xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"])
    for name, body in parts.items():
        assert b"<f>" not in body and b"<f " not in body and b"External" not in body and b"vba" not in body.lower(), name


def test_export_is_readable_by_openpyxl_with_the_same_values() -> None:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(clean()), data_only=False)
    assert wb.sheetnames == ["meta", "allocations"]
    rows = list(wb["allocations"].iter_rows(values_only=True))
    assert rows == [("instrument_id", "target_weight", "note"), ("SPY", 0.6, None), ("CASH", 0.4, None)]
    meta = {r[0]: r[1] for r in wb["meta"].iter_rows(min_row=2, values_only=True)}
    assert meta["base_plan_version_id"] == SAMPLE_META["base_plan_version_id"] and meta["expected_revision"] == 2


def test_workbook_saved_by_openpyxl_imports() -> None:
    """A workbook re-saved by a spreadsheet library (shared strings, docProps, theme) still imports."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(clean()))
    wb["allocations"]["B2"] = 0.5
    wb["allocations"]["B3"] = 0.5
    buf = io.BytesIO()
    wb.save(buf)
    parsed = import_workbook_bytes(buf.getvalue(), LIMITS)
    assert content_from_workbook(SAMPLE_CONTENT, parsed)["allocation"] == {"weights": [{"instrument_id": "SPY", "weight": 0.5}], "cash_weight": 0.5}


# ===================================================================== XLS-02 accepted formats
def test_vba_project_in_xlsx_is_rejected_XLS_02() -> None:
    data = craft(add={"xl/vbaProject.bin": b"\xd0\xcf\x11\xe0synthetic-not-a-real-vba-project"})
    err = rejected(data, "macro_content_rejected")
    assert "macro" in err.message


def test_macro_enabled_content_type_is_rejected_XLS_02() -> None:
    data = craft(replace={"[Content_Types].xml": ("spreadsheetml.sheet.main+xml", "ms-excel.sheet.macroEnabled.main+xml")})
    rejected(data, "macro_content_rejected")


def test_vba_relationship_is_rejected_XLS_02() -> None:
    rel = '<Relationship Id="rId9" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vba.bin"/></Relationships>'
    rejected(craft(replace={"xl/_rels/workbook.xml.rels": ("</Relationships>", rel)}), "macro_content_rejected")


def test_excel4_macro_sheet_is_rejected_XLS_02() -> None:
    rejected(craft(add={"xl/macrosheets/sheet1.xml": sheet_xml([])}), "macro_content_rejected")


@pytest.mark.parametrize("part", ["xl/activeX/activeX1.xml", "xl/embeddings/oleObject1.bin", "xl/activeX/activeX1.bin"])
def test_activex_and_ole_objects_are_rejected_XLS_02(part: str) -> None:
    rejected(craft(add={part: b"<x/>" if part.endswith(".xml") else b"synthetic"}), "embedded_object_rejected")


def test_legacy_and_binary_formats_are_rejected_XLS_02() -> None:
    rejected(bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 512, "legacy_format_rejected")
    rejected(b"not a workbook at all", "unsupported_format")
    xlsb = craft(add={"xl/workbook.bin": b"binary"})
    rejected(xlsb, "unsupported_format")
    template = craft(replace={"[Content_Types].xml": ("spreadsheetml.sheet.main+xml", "spreadsheetml.template.main+xml")})
    rejected(template, "unsupported_format")


def test_oversized_file_is_rejected_before_parsing_XLS_02() -> None:
    limits = {**LIMITS, "excel_upload_max_bytes": 1024}
    rejected(clean(), "file_too_large", limits=limits)


def test_decompression_bomb_is_rejected_before_parsing_XLS_02() -> None:
    bomb = craft(add={"xl/worksheets/padding.xml": b"<a>" + b" " * (4 * 1024 * 1024) + b"</a>"})
    info = next(i for i in zipfile.ZipFile(io.BytesIO(bomb)).infolist() if i.filename == "xl/worksheets/padding.xml")
    assert info.file_size / info.compress_size > float(LIMITS["excel_max_compression_ratio"])
    rejected(bomb, "compression_ratio_exceeded")


def test_total_decompressed_size_limit_XLS_02() -> None:
    limits = {**LIMITS, "excel_max_uncompressed_bytes": 2048}
    big = rezip({**unzip(clean()), "xl/worksheets/extra.xml": b"<a>" + bytes(range(256)) * 16 + b"</a>"}, compression=zipfile.ZIP_STORED)
    rejected(big, "package_too_large", limits=limits)


def test_forged_entry_sizes_cannot_inflate(monkeypatch: pytest.MonkeyPatch) -> None:
    """A directory that under-declares an entry's size is caught by the bounded reader."""
    data = craft(add={"xl/worksheets/pad.xml": b"<a>" + b"x" * 5000 + b"</a>"})
    real_infolist = zipfile.ZipFile.infolist

    def lying_infolist(self: zipfile.ZipFile) -> list[zipfile.ZipInfo]:
        out = real_infolist(self)
        for i in out:
            if i.filename == "xl/worksheets/pad.xml":
                i.file_size = 100  # the forged directory claims 100 bytes
        return out

    monkeypatch.setattr(zipfile.ZipFile, "infolist", lying_infolist)
    with pytest.raises((ImportRejected, zipfile.BadZipFile)):
        import_workbook_bytes(data, LIMITS)


def test_unsafe_part_names_are_rejected() -> None:
    parts = {**unzip(clean()), "../evil.xml": b"<a/>"}
    rejected(rezip(parts), "unsupported_format")


# ===================================================================== XLS-03 no code or formula execution
def _with_alloc_sheet(rows: list[list[str]]) -> bytes:
    return craft(add={"xl/worksheets/sheet2.xml": sheet_xml(rows)})


HEADER = ['<c r="A1" t="inlineStr"><is><t>instrument_id</t></is></c>', '<c r="B1" t="inlineStr"><is><t>target_weight</t></is></c>', '<c r="C1" t="inlineStr"><is><t>note</t></is></c>']


def test_formula_in_allocation_cell_is_rejected_naming_the_cell_XLS_03() -> None:
    data = _with_alloc_sheet(
        [
            HEADER,
            ['<c r="A2" t="inlineStr"><is><t>SPY</t></is></c>', '<c r="B2"><f>WEBSERVICE("https://example.invalid/x")</f><v>0.6</v></c>'],
            ['<c r="A3" t="inlineStr"><is><t>CASH</t></is></c>', '<c r="B3"><v>0.4</v></c>'],
        ]
    )
    err = rejected(data, "workbook_content_rejected")
    f = err.details["findings"][0]
    assert (f["sheet"], f["row"], f["column"], f["details"]["cell"], f["details"]["reason"]) == ("allocations", 2, "B", "B2", "formula_not_allowed")


def test_formula_tolerant_cell_uses_cached_value_only_XLS_03() -> None:
    data = _with_alloc_sheet(
        [
            HEADER,
            ['<c r="A2" t="inlineStr"><is><t>SPY</t></is></c>', '<c r="B2"><f>1-0.4</f><v>0.6</v></c>'],
            ['<c r="A3" t="inlineStr"><is><t>CASH</t></is></c>', '<c r="B3"><v>0.4</v></c>'],
        ]
    )
    spec = TemplateSpec(TEMPLATE_VERSION, formula_tolerant_columns=frozenset({"target_weight"}))
    parsed = import_workbook_bytes(data, LIMITS, spec=spec)
    assert parsed.rows[0] == {"row": 2, "instrument_id": "SPY", "target_weight": 0.6, "formula_tolerant": True}
    # a formula with a different (stale) cached value is NOT recomputed: the cached value is used
    stale = _with_alloc_sheet([HEADER, ['<c r="A2" t="inlineStr"><is><t>SPY</t></is></c>', '<c r="B2"><f>1-0.4</f><v>0.25</v></c>']])
    assert import_workbook_bytes(stale, LIMITS, spec=spec).rows[0]["target_weight"] == 0.25


def test_formula_in_meta_sheet_is_always_rejected_XLS_03() -> None:
    meta = unzip(clean())["xl/worksheets/sheet1.xml"].decode()
    tampered = meta.replace('<c r="B7"><v>2</v></c>', '<c r="B7"><f>1+1</f><v>2</v></c>')
    assert tampered != meta
    rejected(craft(add={"xl/worksheets/sheet1.xml": tampered}), "workbook_content_rejected")


def test_external_workbook_link_is_rejected_without_network_XLS_03(monkeypatch: pytest.MonkeyPatch) -> None:
    attempts: list[Any] = []
    monkeypatch.setattr(socket.socket, "connect", lambda *a, **k: attempts.append(a))
    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **k: attempts.append(a) or [])
    link = '<?xml version="1.0"?><externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><externalBook/></externalLink>'
    rels = '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath" Target="https://example.invalid/book.xlsx" TargetMode="External"/></Relationships>'
    data = craft(add={"xl/externalLinks/externalLink1.xml": link, "xl/externalLinks/_rels/externalLink1.xml.rels": rels})
    rejected(data, "external_link_rejected")
    # an external relationship anywhere else (for example a hyperlink) is rejected too
    hyper = '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.invalid/" TargetMode="External"/></Relationships>'
    rejected(craft(add={"xl/worksheets/_rels/sheet2.xml.rels": hyper}), "external_link_rejected")
    assert attempts == []


def test_xxe_and_entity_expansion_are_rejected_XLS_03() -> None:
    xxe = f'<?xml version="1.0"?><!DOCTYPE w [<!ENTITY x SYSTEM "file:///etc/hostname">]><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&x;</t></is></c></row></sheetData></worksheet>'
    rejected(craft(add={"xl/worksheets/sheet2.xml": xxe}), "xml_entities_rejected")
    laughs = '<?xml version="1.0"?><!DOCTYPE l [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">]><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>&b;</t></si></sst>'
    rejected(craft(add={"xl/sharedStrings.xml": laughs}), "xml_entities_rejected")
    rejected(craft(add={"xl/styles.xml": "<styleSheet"}), "malformed_xml")


def test_no_network_during_a_full_parse_XLS_03(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[str] = []
    for name in ("connect", "connect_ex", "sendto", "send"):
        monkeypatch.setattr(socket.socket, name, lambda *a, _n=name, **k: calls.append(_n))
    import subprocess

    monkeypatch.setattr(subprocess, "Popen", lambda *a, **k: calls.append("subprocess"))
    parsed = import_workbook_bytes(clean(), LIMITS)
    assert parsed.rows and calls == []


# ===================================================================== template content findings
def test_header_and_unknown_meta_keys_are_named() -> None:
    bad_header = craft(replace={"xl/worksheets/sheet2.xml": ("<t xml:space=\"preserve\">target_weight</t>", "<t xml:space=\"preserve\">weight</t>")})
    err = rejected(bad_header, "workbook_content_rejected")
    assert err.details["findings"][0]["row"] == 1 and err.details["findings"][0]["sheet"] == "allocations"
    extra_key = craft(replace={"xl/worksheets/sheet1.xml": ("</sheetData>", '<row r="20"><c r="A20" t="inlineStr"><is><t>macro</t></is></c><c r="B20" t="inlineStr"><is><t>x</t></is></c></row></sheetData>')})
    err = rejected(extra_key, "workbook_content_rejected")
    assert {"sheet": "meta", "row": 20, "column": "A"}.items() <= err.details["findings"][0].items()


def test_non_numeric_weight_and_duplicates_are_named() -> None:
    data = _with_alloc_sheet(
        [
            HEADER,
            ['<c r="A2" t="inlineStr"><is><t>SPY</t></is></c>', '<c r="B2" t="inlineStr"><is><t>sixty</t></is></c>'],
            ['<c r="A3" t="inlineStr"><is><t>CASH</t></is></c>', '<c r="B3"><v>0.4</v></c>'],
            ['<c r="A4" t="inlineStr"><is><t>CASH</t></is></c>', '<c r="B4"><v>0.4</v></c>'],
        ]
    )
    err = rejected(data, "workbook_content_rejected")
    rows = {(f["row"], f.get("field")) for f in err.details["findings"]}
    assert (2, "target_weight") in rows and (4, "instrument_id") in rows


def test_weight_out_of_contract_range_maps_back_to_row() -> None:
    data = _with_alloc_sheet([HEADER, ['<c r="A2" t="inlineStr"><is><t>SPY</t></is></c>', '<c r="B2"><v>1.5</v></c>']])
    err = rejected(data, "workbook_content_rejected")
    assert {"sheet": "allocations", "row": 2, "column": "B"}.items() <= err.details["findings"][0].items()


def test_open_package_returns_parts_for_a_clean_workbook() -> None:
    pkg = open_package(clean(), PackageLimits.from_config(LIMITS))
    assert "xl/workbook.xml" in pkg.parts and pkg.content_types["xl/workbook.xml"].endswith("sheet.main+xml")
